#!/usr/bin/env python3
# /// script
# requires-python = ">=3.10"
# dependencies = ["openpyxl>=3.1"]
# ///
"""Deterministic half of doc-distill: snapshot, diff, glossary subset, and spec checks.

  sync SOURCE --slug S   convert SOURCE to a line snapshot; when it differs from
                         the last checked version, write a diff that maps the
                         changed rows to the spec rules citing them
  glossary --slug S      print only the glossary rows whose source term occurs
                         in the current snapshot
  check --slug S         validate the spec against the current snapshot and, on
                         success, record that version as the new baseline

The baseline lives in docs/specs/<slug>/.distill.json and is committed with the
spec: row hashes of the last checked version, plus each rule's line hash and
refs. The snapshot in .doc-distill/ is a cache that sync can always rebuild.

Run with `uv run distill.py ...` so openpyxl is provided without touching the
project environment. Non-Excel sources are converted with MarkItDown via `uvx`.
"""
import argparse
import difflib
import hashlib
import json
import re
import shutil
import subprocess
import sys
from datetime import date, datetime, time, timedelta
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

EXCEL_SUFFIXES = {".xlsx", ".xlsm"}
MARKITDOWN = ["uvx", "--from", "markitdown[docx,pdf,pptx,xls]", "markitdown"]
RULE_RE = re.compile(r"^\s*[-*]\s+(~~)?\*\*([A-Z]{2,5}-\d{3,})\*\*")
REMOVED_RE = re.compile(r"\[removed \d{4}-\d{2}-\d{2}\]")
CELL_RE = re.compile(r"^([A-Z]{0,3})(\d+)$")
LINE_REF_RE = re.compile(r"^L(\d+)(?:-L?(\d+))?$")


class DistillError(Exception):
    pass


# ---------- paths and stored data ----------

def paths(args) -> dict:
    root = Path(args.cache) / args.slug
    spec_dir = Path(args.specs) / args.slug
    return {
        "root": root,
        "current": root / "current",
        "manifest": root / "manifest.json",
        "report": root / "diff.md",
        "spec": spec_dir / "spec.md",
        "changes": spec_dir / "changes.md",
        "state": spec_dir / ".distill.json",
    }


def load_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8")) if path.exists() else {}


def save_json(path: Path, data: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def file_sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def text_hash(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()[:12]


def encode_rows(entries: list) -> str:
    return " ".join(f"{loc}:{digest}" for loc, digest in entries)


def decode_rows(text: str) -> list:
    return [(int(loc), digest) for loc, _, digest in (item.partition(":") for item in text.split())]


# ---------- cell rendering ----------

def general_number(value) -> str:
    """Excel's General format: whole numbers as is, others to 11 significant digits."""
    if isinstance(value, int) or (float(value).is_integer() and abs(value) < 1e15):
        return str(int(value))
    return f"{value:.11g}"


def render_digits(number: Decimal, pattern: str) -> str:
    """Render a number with a digit pattern such as '#,##0.00'."""
    whole, _, fraction = pattern.partition(".")
    scale = len(whole) - len(whole.rstrip(","))  # each trailing comma divides by 1000
    whole = whole.rstrip(",")
    places = sum(ch in "0#?" for ch in fraction)
    required = fraction.count("0")
    rounded = (number / 1000 ** scale).quantize(Decimal(1).scaleb(-places), rounding=ROUND_HALF_UP)
    text = f"{abs(rounded):{',' if ',' in whole else ''}f}"
    head, dot, tail = text.partition(".")
    if places > required:
        tail = tail[:required] + tail[required:].rstrip("0")
        dot = "." if tail else ""
    if "," not in whole:
        head = head.zfill(whole.count("0"))
    return ("-" if rounded < 0 else "") + head + dot + tail


def format_number(value, number_format: str) -> str:
    """Render a number with the first section of its Excel number format.

    Covers what specs rely on (separators, decimals, percent, currency and unit
    literals); fractions, dates, and scientific formats fall back to General.
    """
    section = number_format.split(";")[0]
    section = re.sub(r"\[\$([^\]-]*)[^\]]*\]", lambda m: f'"{m.group(1)}"', section)  # [$¥-411] -> "¥"
    section = re.sub(r"\[[^\]]*\]", "", section)  # colors, conditions, locales
    parts, digits, percent, i = [], "", False, 0
    while i < len(section):
        ch = section[i]
        if ch == '"':
            end = section.find('"', i + 1)
            end = len(section) if end < 0 else end
            parts.append(section[i + 1:end])
            i = end + 1
        elif ch == "\\":
            parts.append(section[i + 1:i + 2])
            i += 2
        elif ch in "_*":
            i += 2  # padding and fill take the next character and show nothing we keep
        elif section.startswith("General", i):
            parts.append(None)
            digits = "General"
            i += len("General")
        elif ch in "0#?,.":
            if None not in parts:
                parts.append(None)
            digits += ch
            i += 1
        elif ch in "/@" or (ch.isascii() and ch.isalpha()):
            return general_number(value)
        else:
            percent = percent or ch == "%"
            parts.append(ch)
            i += 1
    if None not in parts:
        return general_number(value)
    number = Decimal(repr(value)) if isinstance(value, float) else Decimal(value)
    if percent:
        number *= 100
    rendered = general_number(float(number)) if digits == "General" else render_digits(number, digits)
    return "".join(rendered if part is None else part for part in parts)


def display_value(value, number_format: str) -> str:
    """Render a cell value as Excel shows it, with dates in ISO form."""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime):
        if value.time() == time():
            return value.date().isoformat()
        return value.strftime("%Y-%m-%d %H:%M:%S" if value.second else "%Y-%m-%d %H:%M")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.strftime("%H:%M:%S" if value.second else "%H:%M")
    if isinstance(value, timedelta):
        return str(value)
    if isinstance(value, (int, float)):
        return format_number(value, number_format or "General")
    return str(value)


def cell_text(cell) -> str:
    if cell.value is None:
        return ""
    text = display_value(cell.value, cell.number_format).strip()
    return text.replace("\r\n", "\n").replace("\n", " ↵ ").replace("|", "\\|")


# ---------- conversion ----------

def convert_excel(source: Path, out: Path) -> list:
    from openpyxl import load_workbook

    # Cached formula results are what the author saw. Read-only mode streams the
    # file, and reset_dimensions() ignores the stored used range, so a styled
    # cell far from the content does not make us walk a huge empty grid.
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        units = []
        for index, sheet in enumerate(workbook.worksheets, start=1):
            sheet.reset_dimensions()
            lines = []
            for row in sheet.iter_rows():
                cells = [(cell, text) for cell in row if (text := cell_text(cell))]
                if cells:
                    shown = " | ".join(f"{cell.column_letter}: {text}" for cell, text in cells)
                    lines.append(f"{cells[0][0].row} | {shown}")
            safe_title = re.sub(r"[^\w.-]+", "_", sheet.title)
            file_name = f"{index:02d}-{safe_title}.txt"
            (out / file_name).write_text("".join(f"{line}\n" for line in lines), encoding="utf-8")
            units.append({"name": sheet.title, "file": file_name, "hidden": sheet.sheet_state != "visible"})
        return units
    finally:
        workbook.close()


def convert_other(source: Path, out: Path) -> list:
    try:
        result = subprocess.run([*MARKITDOWN, str(source)], capture_output=True, text=True, check=False)
    except FileNotFoundError as error:
        raise DistillError("uvx not found; install uv to convert non-Excel sources") from error
    if result.returncode != 0:
        raise DistillError(f"markitdown failed: {result.stderr.strip()[-500:]}")
    (out / "doc.txt").write_text(result.stdout, encoding="utf-8")
    return [{"name": "doc", "file": "doc.txt", "hidden": False}]


def kind_of(source: Path) -> str:
    return "excel" if source.suffix.lower() in EXCEL_SUFFIXES else "doc"


def convert(source: Path, sha: str, p: dict) -> dict:
    kind = kind_of(source)
    staging = p["root"] / "staging"
    shutil.rmtree(staging, ignore_errors=True)
    staging.mkdir(parents=True)
    units = convert_excel(source, staging) if kind == "excel" else convert_other(source, staging)
    shutil.rmtree(p["current"], ignore_errors=True)
    staging.rename(p["current"])
    manifest = {"source": str(source), "source_sha": sha, "kind": kind,
                "synced": date.today().isoformat(), "units": units}
    save_json(p["manifest"], manifest)
    return manifest


def load_unit(path: Path, kind: str) -> list:
    """Return [(locator, content)] for non-blank lines: the Excel row or the doc line number."""
    entries = []
    for number, line in enumerate(path.read_text(encoding="utf-8").splitlines(), start=1):
        if not line.strip():
            continue
        if kind == "doc":
            entries.append((number, line))
        else:
            head, _, rest = line.partition(" | ")
            entries.append((int(head), rest))
    return entries


def current_units(p: dict, manifest: dict) -> dict:
    return {u["name"]: load_unit(p["current"] / u["file"], manifest["kind"]) for u in manifest["units"]}


def hashed(units: dict) -> dict:
    return {name: [(loc, text_hash(content)) for loc, content in entries] for name, entries in units.items()}


# ---------- references ----------

def split_refs(line: str) -> list:
    """Return the refs in every (src: ...) group, honoring quoted names and nested parentheses."""
    refs = []
    for match in re.finditer(r"\(src:", line):
        depth, quoted, part = 0, False, ""
        for ch in line[match.end():]:
            if ch == "'":
                quoted = not quoted
            elif not quoted and ch == "(":
                depth += 1
            elif not quoted and ch == ")":
                if depth == 0:
                    break
                depth -= 1
            elif not quoted and ch == "," and depth == 0:
                refs.append(part)
                part = ""
                continue
            part += ch
        refs.append(part)
    return [ref.strip() for ref in refs if ref.strip()]


def parse_ref(ref: str, kind: str):
    """Return (unit, first row, last row, first column, last column), or None when malformed."""
    if kind == "doc":
        match = LINE_REF_RE.match(ref)
        if not match:
            return None
        first, last = int(match.group(1)), int(match.group(2) or match.group(1))
        return "doc", min(first, last), max(first, last), "", ""
    sheet, sep, cells = ref.rpartition("!")
    if not sep or not sheet:
        return None
    if len(sheet) > 1 and sheet[0] == sheet[-1] == "'":
        sheet = sheet[1:-1].replace("''", "'")
    bounds = [CELL_RE.match(part) for part in cells.split(":")]
    if not all(bounds) or len(bounds) > 2:
        return None
    rows = (int(bounds[0].group(2)), int(bounds[-1].group(2)))
    return sheet, min(rows), max(rows), bounds[0].group(1), bounds[-1].group(1)


def format_ref(unit: str, first: int, last: int, first_col: str, last_col: str, kind: str) -> str:
    if kind == "doc":
        return f"L{first}" if first == last else f"L{first}-L{last}"
    sheet = unit if re.fullmatch(r"\w+", unit) else "'" + unit.replace("'", "''") + "'"
    if first == last and first_col == last_col:
        return f"{sheet}!{first_col}{first}"
    return f"{sheet}!{first_col}{first}:{last_col}{last}"


def parse_rules(spec: Path) -> list:
    rules = []
    for number, line in enumerate(spec.read_text(encoding="utf-8").splitlines(), start=1):
        match = RULE_RE.match(line)
        if match:
            removed = bool(match.group(1)) or bool(REMOVED_RE.search(line))
            rules.append({"id": match.group(2), "number": number, "text": line.strip(),
                          "refs": split_refs(line), "removed": removed})
    return rules


def cited(units: dict, refs: list, kind: str) -> list:
    """Return the (locator, content) entries that refs cover, in ref order."""
    entries = []
    for ref in refs:
        parsed = parse_ref(ref, kind)
        if parsed and parsed[0] in units:
            unit, first, last = parsed[:3]
            entries += [entry for entry in units[unit] if first <= entry[0] <= last]
    return entries


# ---------- sync ----------

def map_rows(first: int, last: int, old: list, new: list, codes: list):
    """Follow old rows first..last into the new version.

    Returns (status, first, last) with status "same", "moved", "changed", or "removed".
    """
    inside = [loc for loc, _ in old if first <= loc <= last]
    mapped, related, changed = {}, [], False
    for tag, i1, i2, j1, j2 in codes:
        hits = [i for i in range(i1, i2) if first <= old[i][0] <= last]
        if tag == "equal":
            for i in hits:
                mapped[old[i][0]] = new[j1 + i - i1][0]
        elif tag == "insert":
            if 0 < i1 < len(old) and first <= old[i1 - 1][0] and old[i1][0] <= last:
                changed = True
                related += [new[j][0] for j in range(j1, j2)]
        elif hits:
            changed = True
            related += [new[j][0] for j in range(j1, j2)]
    if not changed and inside:
        lo = first + mapped[inside[0]] - inside[0]
        hi = last + mapped[inside[-1]] - inside[-1]
        if [loc for loc, _ in new if lo <= loc <= hi] != [mapped[loc] for loc in inside]:
            return "changed", lo, hi  # new content appeared in rows the ref spans
        return ("same" if (lo, hi) == (first, last) else "moved"), lo, hi
    related += mapped.values()
    if not related:
        return "removed", first, last
    return "changed", min(related), max(related)


def follow_ref(ref: str, kind: str, old: dict, new: dict, codes: dict):
    """Follow one baseline ref into the new version: (status, new ref or None)."""
    parsed = parse_ref(ref, kind)
    if not parsed or parsed[0] not in codes:
        return "removed", None
    unit, first, last, first_col, last_col = parsed
    status, lo, hi = map_rows(first, last, old[unit], new[unit], codes[unit])
    if status == "removed":
        return status, None
    return status, format_ref(unit, lo, hi, first_col, last_col, kind)


def show(entry, kind: str) -> str:
    loc, content = entry
    return f"L{loc}: {content}" if kind == "doc" else f"{loc} | {content}"


def spans(locators: list, kind: str) -> str:
    """Compress [3, 4, 5, 9] into 'rows 3-5, 9' ('lines L3-L5, L9' for documents)."""
    prefix = "L" if kind == "doc" else ""
    groups = []
    for loc in locators:
        if groups and loc == groups[-1][1] + 1:
            groups[-1][1] = loc
        else:
            groups.append([loc, loc])
    text = ", ".join(f"{prefix}{a}" if a == b else f"{prefix}{a}-{prefix}{b}" for a, b in groups)
    noun = "line" if kind == "doc" else "row"
    return f"{noun}{'s' if len(locators) > 1 else ''} {text}"


def write_diff(p: dict, manifest: dict, state: dict, slug: str) -> dict:
    """Diff the current snapshot against the checked baseline and write diff.md."""
    kind = manifest["kind"]
    current = current_units(p, manifest)
    new = hashed(current)
    old = {name: decode_rows(rows) for name, rows in state["units"].items()}
    codes = {
        name: difflib.SequenceMatcher(a=[h for _, h in old[name]], b=[h for _, h in new[name]],
                                      autojunk=False).get_opcodes()
        for name in new if name in old
    }

    changes, changed_units = [], []
    for name, entries in current.items():
        if name not in old:
            heading, lines = f"### {name} (sheet added)", [f"+ {show(entry, kind)}" for entry in entries] or ["(empty)"]
        else:
            heading, lines = f"### {name}", []
            for tag, i1, i2, j1, j2 in codes[name]:
                if tag == "equal":
                    continue
                gone = spans([loc for loc, _ in old[name][i1:i2]], kind)
                lines.append({"insert": "- inserted:", "delete": f"- removed old {gone}",
                              "replace": f"- replaced old {gone} with:"}[tag])
                lines += [f"  + {show(entries[j], kind)}" for j in range(j1, j2)]
        if lines:
            changed_units.append(name)
            changes += [heading, "", *lines, ""]
    for name in old.keys() - current.keys():
        changed_units.append(name)
        changes += [f"### {name} (sheet removed)", ""]

    spec_text = {rule["id"]: rule["text"] for rule in parse_rules(p["spec"])} if p["spec"].exists() else {}
    affected_ids, affected, moved = [], [], []
    for rule_id, info in state["rules"].items():
        if info.get("removed"):
            continue
        outcomes = [(ref, *follow_ref(ref, kind, old, new, codes)) for ref in info["refs"]]
        if any(status in ("changed", "removed") for _, status, _ in outcomes):
            now = ", ".join(target or f"{ref} gone" for ref, _, target in outcomes)
            affected_ids.append(rule_id)
            affected += [f"- {rule_id} (now {now})", f"  {spec_text.get(rule_id, '(not in the spec)')}"]
        else:
            moved += [f"- {rule_id}: {ref} -> {target}" for ref, status, target in outcomes if status == "moved"]

    report = [f"# Diff for {slug}: sha256 {state['source_sha'][:8]} -> {manifest['source_sha'][:8]}", ""]
    if affected:
        report += ["## Affected rules", "", "The rows they cite changed. Update each rule and its ref.", "",
                   *affected, ""]
    if moved:
        report += ["## Moved references", "", "Content unchanged; update only the ref.", "", *moved, ""]
    report += ["## Changes", "", *(changes or ["No content changes (formatting or metadata only).", ""])]
    p["report"].write_text("\n".join(report), encoding="utf-8")
    return {"units": changed_units, "affected": affected_ids, "moved": len(moved)}


def sync(args) -> int:
    source = Path(args.source)
    if not source.is_file():
        raise DistillError(f"source not found: {source}")
    p = paths(args)
    manifest = load_json(p["manifest"])
    sha = file_sha(source)
    if args.force or manifest.get("source_sha") != sha or not p["current"].exists():
        manifest = convert(source, sha, p)
    state = load_json(p["state"])
    label = f"{source.name} (sha256 {sha[:8]})"

    if not state:
        listing = ", ".join(f"{u['file']}{' (hidden)' if u['hidden'] else ''}" for u in manifest["units"])
        if p["spec"].exists():
            print(f"snapshot: {label} -> {p['current']}/ [{listing}]")
            print(f"warning: {p['spec']} exists but {p['state']} does not; make sure its refs match "
                  "this version, then run check to record it as the baseline")
        else:
            print(f"new: {label} -> {p['current']}/ [{listing}]")
            print("next: write the spec from the snapshot files")
        return 0
    if state["source_sha"] == sha:
        print(f"unchanged: {label} is the last checked version")
        return 0

    summary = write_diff(p, manifest, state, args.slug)
    print(f"changed: {label}, last checked {state['source_sha'][:8]}; report {p['report']}")
    print(f"changed units: {', '.join(summary['units']) or 'none'}")
    print(f"affected rules: {', '.join(summary['affected']) or 'none'}")
    print(f"moved refs: {summary['moved']}")
    return 0


# ---------- glossary ----------

def glossary(args) -> int:
    p = paths(args)
    manifest = load_json(p["manifest"])
    if not manifest:
        raise DistillError(f"no snapshot for {args.slug}; run sync first")
    text = "\n".join((p["current"] / u["file"]).read_text(encoding="utf-8") for u in manifest["units"])
    path = Path(args.glossary)
    if not path.exists():
        print(f"no glossary at {path}")
        return 0
    rows = [line for line in path.read_text(encoding="utf-8").splitlines() if line.startswith("|")]
    if len(rows) < 2:
        print("glossary has no table")
        return 0
    print("\n".join(rows[:2]))
    hits = 0
    for row in rows[2:]:
        term = row.strip("|").split("|")[0].strip()
        if term and term in text:
            print(row)
            hits += 1
    print(f"\n{hits} of {len(rows) - 2} terms occur in the source", file=sys.stderr)
    return 0


# ---------- check ----------

def check(args) -> int:
    p = paths(args)
    manifest = load_json(p["manifest"])
    if not manifest:
        raise DistillError(f"no snapshot for {args.slug}; run sync first")
    if not p["spec"].exists():
        raise DistillError(f"spec not found: {p['spec']}")
    kind = manifest["kind"]
    units = hashed(current_units(p, manifest))
    state = load_json(p["state"])
    baseline = {name: decode_rows(rows) for name, rows in state.get("units", {}).items()}
    known = state.get("rules", {})

    errors, seen = [], {}
    rules = parse_rules(p["spec"])
    for rule in rules:
        where = f"{p['spec']}:{rule['number']} {rule['id']}"
        if rule["id"] in seen:
            errors.append(f"{where}: duplicate ID (first at line {seen[rule['id']]})")
        seen.setdefault(rule["id"], rule["number"])
        if rule["removed"]:
            continue
        if not rule["refs"]:
            errors.append(f"{where}: no (src: ...) reference")
        for ref in rule["refs"]:
            parsed = parse_ref(ref, kind)
            if not parsed:
                errors.append(f"{where}: malformed ref '{ref}'")
            elif parsed[0] not in units:
                errors.append(f"{where}: unknown sheet '{parsed[0]}' in '{ref}'")
            elif not cited(units, [ref], kind):
                errors.append(f"{where}: '{ref}' points at empty or missing rows")
        previous = known.get(rule["id"], {})
        if (previous.get("line") == text_hash(rule["text"]) and rule["id"] not in args.accept
                and cited(baseline, previous["refs"], kind) != cited(units, rule["refs"], kind)):
            errors.append(f"{where}: the rows it cites changed since the last check but the rule was not "
                          f"edited; update its text or ref, or pass --accept {rule['id']} if it still holds")

    for lost in sorted(known.keys() - seen.keys()):
        errors.append(f"{lost}: ID disappeared from the spec; keep it and mark it [removed YYYY-MM-DD]")

    short = manifest["source_sha"][:8]
    if state.get("source_sha") != manifest["source_sha"]:
        log = p["changes"].read_text(encoding="utf-8") if p["changes"].exists() else ""
        if f"sha256 {short}" not in log:
            errors.append(f"{p['changes']}: no entry for this version (sha256 {short})")

    if errors:
        print("\n".join(errors))
        print(f"\n{len(errors)} problem(s) in {len(rules)} rules")
        return 1
    save_json(p["state"], {
        "source_sha": manifest["source_sha"],
        "units": {name: encode_rows(entries) for name, entries in units.items()},
        "rules": {
            rule["id"]: {"removed": True} if rule["removed"]
            else {"line": text_hash(rule["text"]), "refs": rule["refs"]}
            for rule in rules
        },
    })
    print(f"ok: {len(rules)} rules; baseline is now sha256 {short}")
    return 0


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    sub = parser.add_subparsers(dest="command", required=True)
    sync_parser = sub.add_parser("sync")
    sync_parser.add_argument("source")
    sync_parser.add_argument("--force", action="store_true", help="re-convert even when the hash is unchanged")
    glossary_parser = sub.add_parser("glossary")
    glossary_parser.add_argument("--glossary", default="docs/glossary.md")
    check_parser = sub.add_parser("check")
    check_parser.add_argument("--accept", nargs="+", default=[], metavar="ID",
                              help="rules whose cited rows changed but which still hold")
    for command in sub.choices.values():
        command.add_argument("--slug", required=True, help="stable name for this document, e.g. login-spec")
        command.add_argument("--cache", default=".doc-distill", help="snapshot cache (keep it out of git)")
        command.add_argument("--specs", default="docs/specs", help="directory holding <slug>/spec.md")
    args = parser.parse_args()
    try:
        return {"sync": sync, "glossary": glossary, "check": check}[args.command](args)
    except DistillError as error:
        print(f"error: {error}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    sys.exit(main())
