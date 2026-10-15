# /// script
# requires-python = ">=3.10"
# dependencies = ["openpyxl>=3.1"]
# ///
"""Regression tests for distill.py, run against generated workbooks.

Run: uv run skills-dev/doc-distill/tests/test_distill.py
"""
import hashlib
import shutil
import subprocess
import sys
import tempfile
import time
import unittest
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

SCRIPT = Path(__file__).resolve().parents[1] / "scripts" / "distill.py"


def rows(*texts, start=5, column="B"):
    return {f"{column}{start + offset}": text for offset, text in enumerate(texts)}


class DistillCase(unittest.TestCase):
    def setUp(self):
        self.dir = Path(tempfile.mkdtemp(prefix="distill-test-"))
        self.spec = self.dir / "docs" / "specs" / "s" / "spec.md"
        self.spec.parent.mkdir(parents=True)

    def tearDown(self):
        shutil.rmtree(self.dir)

    def run_cli(self, command, *args):
        return subprocess.run(
            [sys.executable, str(SCRIPT), command, *args, "--slug", "s"],
            cwd=self.dir, capture_output=True, text=True, timeout=120,
        )

    def sync(self, name, *flags):
        return self.run_cli("sync", name, *flags)

    def check(self, *flags):
        return self.run_cli("check", *flags)

    def workbook(self, name, sheets, formats=None):
        book = Workbook()
        book.remove(book.active)
        for title, cells in sheets.items():
            sheet = book.create_sheet(title)
            for ref, value in cells.items():
                sheet[ref] = value
        for (title, ref), number_format in (formats or {}).items():
            book[title][ref].number_format = number_format
        book.save(self.dir / name)
        return name

    def short_hash(self, name):
        return hashlib.sha256((self.dir / name).read_bytes()).hexdigest()[:8]

    def write_spec(self, *lines):
        self.spec.write_text("# Spec\n\n" + "\n".join(lines) + "\n", encoding="utf-8")

    def log_version(self, name):
        with (self.spec.parent / "changes.md").open("a", encoding="utf-8") as log:
            log.write(f"## 2026-10-15 · sha256 {self.short_hash(name)}\n- Updated\n\n")

    def establish(self, name, *spec_lines):
        """Sync a version, write its spec and changes entry, and pass check."""
        self.assertEqual(self.sync(name).returncode, 0)
        self.write_spec(*spec_lines)
        self.log_version(name)
        result = self.check()
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)

    def diff_report(self):
        return (self.dir / ".doc-distill" / "s" / "diff.md").read_text(encoding="utf-8")


class SnapshotTest(DistillCase):
    def test_empty_sheet_does_not_break_check_or_resync(self):
        v1 = self.workbook("v1.xlsx", {"Main": rows("A rule"), "Empty": {}})
        self.establish(v1, "- **BR-001** A (src: Main!B5)")
        v2 = self.workbook("v2.xlsx", {"Main": rows("A rule", "B rule"), "Empty": {}})
        result = self.sync(v2)
        self.assertEqual(result.returncode, 0, result.stderr)

    def test_values_render_with_their_number_format(self):
        cells = {
            "B2": 0.15, "B3": 0.155, "B4": 1234567, "B5": 1234.5,
            "B6": datetime(2026, 10, 1), "B7": 0.1 + 0.2, "B8": 12, "B9": True,
        }
        formats = {
            ("Main", "B2"): "0%", ("Main", "B3"): "0.0%", ("Main", "B4"): '#,##0"円"',
            ("Main", "B5"): "[$¥-411]#,##0", ("Main", "B6"): "yyyy/mm/dd", ("Main", "B8"): '0"件"',
        }
        self.sync(self.workbook("f.xlsx", {"Main": cells}, formats))
        snapshot = (self.dir / ".doc-distill" / "s" / "current" / "01-Main.txt").read_text(encoding="utf-8")
        self.assertEqual(snapshot.splitlines(), [
            "2 | B: 15%", "3 | B: 15.5%", "4 | B: 1,234,567円", "5 | B: ¥1,235",
            "6 | B: 2026-10-01", "7 | B: 0.3", "8 | B: 12件", "9 | B: TRUE",
        ])

    def test_far_styled_cell_does_not_expand_the_grid(self):
        book = Workbook()
        book.active.title = "Grid"
        book.active["A1"] = "title"
        book.active.cell(2000, 2000).font = Font(bold=True)
        book.save(self.dir / "grid.xlsx")
        started = time.monotonic()
        result = self.sync("grid.xlsx")
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertLess(time.monotonic() - started, 5)


class TrackingTest(DistillCase):
    def test_every_sync_result_names_the_version_hash(self):
        v1 = self.workbook("v1.xlsx", {"Main": rows("A rule")})
        self.assertIn(f"sha256 {self.short_hash(v1)}", self.sync(v1).stdout)
        self.write_spec("- **BR-001** A (src: Main!B5)")
        self.log_version(v1)
        self.assertEqual(self.check().returncode, 0)
        unchanged = self.sync(v1).stdout
        self.assertIn("unchanged", unchanged)
        self.assertIn(f"sha256 {self.short_hash(v1)}", unchanged)
        v2 = self.workbook("v2.xlsx", {"Main": rows("A rule", "B rule")})
        self.assertIn(f"sha256 {self.short_hash(v2)}", self.sync(v2).stdout)

    def test_row_inserted_inside_cited_range_marks_rule_affected(self):
        v1 = self.workbook("v1.xlsx", {"Main": rows("A", "B", "C", "D")})
        self.establish(v1, "- **BR-001** A to D (src: Main!B5:B8)")
        v2 = self.workbook("v2.xlsx", {"Main": rows("A", "B", "NEW", "C", "D")})
        result = self.sync(v2)
        self.assertIn("affected rules: BR-001", result.stdout)
        self.assertIn("moved refs: 0", result.stdout)

    def test_content_typed_into_a_blank_row_at_the_end_of_a_cited_range_marks_rule_affected(self):
        v1 = self.workbook("v1.xlsx", {"Main": {"B5": "A", "B9": "outside"}})
        self.establish(v1, "- **BR-001** A block (src: Main!B5:B7)")
        v2 = self.workbook("v2.xlsx", {"Main": {"B5": "A", "B7": "NEW", "B9": "outside"}})
        self.assertIn("affected rules: BR-001", self.sync(v2).stdout)

    def test_edited_cited_row_marks_rule_affected_and_shows_new_text(self):
        v1 = self.workbook("v1.xlsx", {"Main": rows("min 8", "other")})
        self.establish(v1, "- **VR-001** Password min 8 (src: Main!B5)", "- **BR-001** other (src: Main!B6)")
        v2 = self.workbook("v2.xlsx", {"Main": rows("min 12", "other")})
        self.assertIn("affected rules: VR-001\n", self.sync(v2).stdout)
        report = self.diff_report()
        self.assertIn("- VR-001 (now Main!B5)", report)
        self.assertIn("5 | B: min 12", report)

    def test_renamed_sheet_marks_its_rules_affected(self):
        v1 = self.workbook("v1.xlsx", {"Main": rows("A")})
        self.establish(v1, "- **BR-001** A (src: Main!B5)")
        v2 = self.workbook("v2.xlsx", {"Login": rows("A")})
        self.assertIn("affected rules: BR-001", self.sync(v2).stdout)
        self.assertIn("- BR-001 (now Main!B5 gone)", self.diff_report())

    def test_moved_ref_keeps_its_column_span(self):
        v1 = self.workbook("v1.xlsx", {"Main": {"B5": "Email", "F5": "max 255"}})
        self.establish(v1, "- **VR-001** Email max 255 (src: Main!B5:F5)")
        v2 = self.workbook("v2.xlsx", {"Main": {"B5": "Intro", "B6": "Email", "F6": "max 255"}})
        self.sync(v2)
        self.assertIn("- VR-001: Main!B5:F5 -> Main!B6:F6", self.diff_report())

    def test_check_rejects_untouched_rules_after_rows_shift(self):
        v1 = self.workbook("v1.xlsx", {"Main": rows("A rule", "B rule", "C rule")})
        self.establish(v1, "- **BR-001** A (src: Main!B5)", "- **BR-002** C (src: Main!B7)")
        v2 = self.workbook("v2.xlsx", {"Main": rows("NEW rule", "A rule", "B rule", "C rule")})
        self.sync(v2)
        self.log_version(v2)
        result = self.check()
        self.assertEqual(result.returncode, 1)
        self.assertIn("BR-001", result.stdout)
        self.assertIn("BR-002", result.stdout)
        self.write_spec("- **BR-001** A (src: Main!B6)", "- **BR-002** C (src: Main!B8)")
        self.assertEqual(self.check().returncode, 0)

    def test_accept_keeps_an_untouched_rule_whose_source_changed(self):
        v1 = self.workbook("v1.xlsx", {"Main": rows("Lock after 5 failures")})
        self.establish(v1, "- **BR-001** Lock after 5 failures (src: Main!B5)")
        v2 = self.workbook("v2.xlsx", {"Main": rows("Lock after 5 failures.")})
        self.sync(v2)
        self.log_version(v2)
        self.assertEqual(self.check().returncode, 1)
        self.assertEqual(self.check("--accept", "BR-001").returncode, 0)
        self.assertEqual(self.check().returncode, 0)

    def test_baseline_stays_until_check_passes(self):
        v1 = self.workbook("v1.xlsx", {"Main": rows("A")})
        self.establish(v1, "- **BR-001** A (src: Main!B5)")
        v2 = self.workbook("v2.xlsx", {"Main": rows("X", "A")})
        self.sync(v2)
        self.sync(v2, "--force")
        self.assertIn("- BR-001: Main!B5 -> Main!B6", self.diff_report())
        v3 = self.workbook("v3.xlsx", {"Main": rows("X", "Y", "A")})
        self.sync(v3)
        self.assertIn("- BR-001: Main!B5 -> Main!B7", self.diff_report())

    def test_added_sheet_content_is_in_the_diff(self):
        v1 = self.workbook("v1.xlsx", {"Main": rows("A")})
        self.establish(v1, "- **BR-001** A (src: Main!B5)")
        v2 = self.workbook("v2.xlsx", {"Main": rows("A"), "Validation": {"B2": "最大10文字"}})
        self.sync(v2)
        self.assertIn("最大10文字", self.diff_report())

    def test_committed_state_survives_a_lost_cache(self):
        v1 = self.workbook("v1.xlsx", {"Main": rows("A", "B")})
        self.establish(v1, "- **BR-001** A (src: Main!B5)", "- **BR-002** B (src: Main!B6)")
        shutil.rmtree(self.dir / ".doc-distill")
        self.assertIn("unchanged", self.sync(v1).stdout)
        self.write_spec("- **BR-001** A (src: Main!B5)")
        result = self.check()
        self.assertEqual(result.returncode, 1)
        self.assertIn("BR-002", result.stdout)


class SpecCheckTest(DistillCase):
    def test_sheet_names_with_parentheses_and_commas_are_valid_refs(self):
        v1 = self.workbook("v1.xlsx", {"画面(1)": {"B2": "Login"}, "Login, Logout": {"B2": "Logout"}})
        self.sync(v1)
        self.write_spec("- **UI-001** Login and logout (src: 画面(1)!B2, 'Login, Logout'!B2)")
        self.log_version(v1)
        result = self.check()
        self.assertEqual(result.returncode, 0, result.stdout)

    def test_check_requires_a_changes_entry_for_the_version(self):
        v1 = self.workbook("v1.xlsx", {"Main": rows("A")})
        self.sync(v1)
        self.write_spec("- **BR-001** A (src: Main!B5)")
        result = self.check()
        self.assertEqual(result.returncode, 1)
        self.assertIn("changes.md", result.stdout)
        self.log_version(v1)
        self.assertEqual(self.check().returncode, 0)

    def test_removed_rule_needs_no_source_but_keeps_its_id(self):
        v1 = self.workbook("v1.xlsx", {"Main": rows("A", "B")})
        self.establish(v1, "- **BR-001** A (src: Main!B5)", "- **BR-002** B (src: Main!B6)")
        self.write_spec("- **BR-001** A (src: Main!B5)", "- ~~**BR-002**~~ B [removed 2026-10-15]")
        result = self.check()
        self.assertEqual(result.returncode, 0, result.stdout)
        self.write_spec("- **BR-001** A (src: Main!B5)")
        result = self.check()
        self.assertEqual(result.returncode, 1)
        self.assertIn("BR-002: ID disappeared", result.stdout)

    def test_check_reports_structural_errors(self):
        v1 = self.workbook("v1.xlsx", {"Main": rows("A")})
        self.sync(v1)
        self.log_version(v1)
        self.write_spec(
            "- **BR-001** A (src: Main!B5)",
            "- **BR-001** duplicate (src: Main!B5)",
            "- **BR-002** no source",
            "- **BR-003** empty row (src: Main!B99)",
        )
        result = self.check()
        self.assertEqual(result.returncode, 1)
        for expected in ("duplicate ID", "BR-002: no (src: ...)", "BR-003: 'Main!B99' points at empty"):
            self.assertIn(expected, result.stdout)


@unittest.skipUnless(shutil.which("uvx"), "uvx converts non-Excel sources")
class DocSourceTest(DistillCase):
    def test_line_refs_follow_inserted_paragraphs(self):
        page = self.dir / "page.html"
        page.write_text("<h1>Spec</h1><p>Users must log in.</p><ul><li>Lock after 5 failures</li></ul>")
        self.establish("page.html", "- **BR-001** Lock after 5 failures (src: L5)")
        page.write_text("<h1>Spec</h1><p>Intro</p><p>Users must log in.</p><ul><li>Lock after 5 failures</li></ul>")
        self.sync("page.html")
        self.assertIn("- BR-001: L5 -> L7", self.diff_report())


if __name__ == "__main__":
    unittest.main()
